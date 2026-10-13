"""Reportes: PDF (reportlab) y Excel (openpyxl)."""
from __future__ import annotations

import io
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (Paragraph, SimpleDocTemplate, Spacer, Table,
                                TableStyle)

AZUL = colors.HexColor("#1e3a5f")
VERDE = colors.HexColor("#16a34a")
GRIS = colors.HexColor("#e5e7eb")


def build_excel(kpis: dict, operaciones: pd.DataFrame, meta: dict,
                predicciones: list[dict]) -> bytes:
    wb = Workbook()

    ws = wb.active
    ws.title = "Resumen"
    ws["A1"] = "CANASAN — Analítica predictiva del cierre de ventas"
    ws["A1"].font = Font(bold=True, size=14, color="1E3A5F")
    ws["A2"] = f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    filas = [
        ("Total de operaciones", kpis["total_ops"]),
        ("Cerradas", kpis["cerradas"]),
        ("Anuladas", kpis["anuladas"]),
        ("En proceso", kpis["en_proceso"]),
        ("TECV conservadora (%)", kpis["tecv_conservadora"]),
        ("TECV sobre resueltas (%)", kpis["tecv_resueltas"]),
        ("TPV global (días)", kpis["tpv_global"]),
        ("Valor del pipeline (S/)", round(kpis["valor_pipeline"], 2)),
        ("Ventas proyectadas (S/)", round(kpis["ventas_proyectadas"], 2)),
        ("Mejor modelo", meta["mejor"]),
        ("AUC (CV 5-fold)", meta["tabla"][0]["AUC"]),
        ("Registros modelo", meta["n_registros"]),
    ]
    ws.append([])
    ws.append(["Indicador", "Valor"])
    for c in ws[4]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1E3A5F")
    for k, v in filas:
        ws.append([k, v])
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 22

    ws2 = wb.create_sheet("Operaciones")
    cols = ["proyecto", "dpto", "cliente", "fecha_separacion", "fecha_cierre",
            "estado", "monto_total", "n_pagos", "dias_venta", "fuente"]
    ws2.append([c.replace("_", " ").upper() for c in cols])
    for c in ws2[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1E3A5F")
    for _, r in operaciones.iterrows():
        ws2.append([
            r.get("proyecto"), r.get("dpto"), r.get("cliente"),
            r.get("fecha_separacion")[:10] if isinstance(r.get("fecha_separacion"), str) else r.get("fecha_separacion"),
            r.get("fecha_cierre")[:10] if isinstance(r.get("fecha_cierre"), str) else r.get("fecha_cierre"),
            r.get("estado"), r.get("monto_total"), r.get("n_pagos"),
            r.get("dias_venta"), r.get("fuente"),
        ])
    for i, w in enumerate([14, 8, 40, 16, 14, 12, 14, 10, 12, 10], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    ws3 = wb.create_sheet("Modelo")
    ws3.append(["Modelo", "Accuracy", "Precision", "Recall", "F1", "AUC"])
    for c in ws3[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1E3A5F")
    for m in meta["tabla"]:
        ws3.append([m["Modelo"], m["Accuracy"], m["Precision"], m["Recall"], m["F1"], m["AUC"]])
    ws3.append([])
    ws3.append(["Mejor modelo (por AUC)", meta["mejor"]])
    ws3.append(["Umbral clasificación ALTA", meta["umbral_alta"]])
    ws3.append(["Umbral clasificación MEDIA", meta["umbral_media"]])
    ws3.append([])
    ws3.append(["Factores del mejor modelo", "Peso (%)"])
    for f in meta["importancias"]:
        ws3.append([f["factor"], f["peso_pct"]])
    ws3.column_dimensions["A"].width = 28
    ws3.column_dimensions["B"].width = 12

    ws4 = wb.create_sheet("Predicciones")
    ws4.append(["Fecha", "Proyecto", "Monto (S/)", "Pagos", "Mes", "Año", "Prob.", "Clase"])
    for c in ws4[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1E3A5F")
    for p in predicciones:
        ws4.append([
            p.get("created_at"), p.get("proyecto"), p.get("monto_total"),
            p.get("n_pagos"), p.get("mes_sep"), p.get("anio_sep"),
            p.get("probabilidad"), p.get("prediccion"),
        ])
    for i, w in enumerate([20, 14, 14, 8, 8, 8, 10, 10], 1):
        ws4.column_dimensions[get_column_letter(i)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_pdf(kpis: dict, meta: dict, top_pipeline: pd.DataFrame,
              factores: list[dict]) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter,
                            leftMargin=2 * cm, rightMargin=2 * cm,
                            topMargin=2 * cm, bottomMargin=2 * cm)
    estilos = getSampleStyleSheet()
    h1 = ParagraphStyle("H1", parent=estilos["Heading1"], textColor=AZUL, fontSize=16, spaceAfter=6)
    h2 = ParagraphStyle("H2", parent=estilos["Heading2"], textColor=AZUL, fontSize=12, spaceBefore=12, spaceAfter=6)
    body = estilos["BodyText"]
    body.fontSize = 10

    elems = [
        Paragraph("CANASAN — Informe de analítica predictiva del cierre de ventas", h1),
        Paragraph(f"Generado el {datetime.now().strftime('%d/%m/%Y %H:%M')} · Surquillo 2026", body),
        Spacer(1, 12),
        Paragraph("1. Indicadores generales", h2),
    ]

    t1 = Table([
        ["Indicador", "Valor"],
        ["Total de operaciones", str(kpis["total_ops"])],
        ["Cerradas", str(kpis["cerradas"])],
        ["Anuladas", str(kpis["anuladas"])],
        ["En proceso", str(kpis["en_proceso"])],
        ["TECV conservadora", f"{kpis['tecv_conservadora']} %"],
        ["TECV sobre resueltas", f"{kpis['tecv_resueltas']} %"],
        ["TPV global", f"{kpis['tpv_global']} días"],
        ["Valor del pipeline", f"S/ {kpis['valor_pipeline']:,.2f}"],
        ["Ventas proyectadas", f"S/ {kpis['ventas_proyectadas']:,.2f}"],
    ], colWidths=[9 * cm, 6 * cm])
    t1.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), AZUL),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("GRID", (0, 0), (-1, -1), 0.4, GRIS),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f3f6fa")]),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    elems += [t1, Paragraph("2. Comparación de modelos (CV estratificado 5-fold)", h2)]

    datos = [["Modelo", "Accuracy", "Precision", "Recall", "F1", "AUC"]]
    for m in meta["tabla"]:
        datos.append([m["Modelo"], f"{m['Accuracy']:.4f}", f"{m['Precision']:.4f}",
                      f"{m['Recall']:.4f}", f"{m['F1']:.4f}", f"{m['AUC']:.4f}"])
    t2 = Table(datos, colWidths=[4.2 * cm] + [2.1 * cm] * 5)
    t2.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), AZUL),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.4, GRIS),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("BACKGROUND", (0, 1), (-1, 1), colors.HexColor("#dcfce7")),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    elems += [t2,
              Paragraph(f"Mejor modelo seleccionado por AUC: <b>{meta['mejor']}</b> "
                        f"(holdout AUC {meta['holdout']['auc']:.4f}, accuracy {meta['holdout']['accuracy']:.4f}).", body),
              Paragraph("3. Factores que más influyen en el cierre", h2)]

    datos3 = [["Factor", "Peso (%)"]] + [[f["factor"], f"{f['peso_pct']} %"] for f in factores[:6]]
    t3 = Table(datos3, colWidths=[9 * cm, 4 * cm])
    t3.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), AZUL),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("GRID", (0, 0), (-1, -1), 0.4, GRIS),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    elems += [t3, Paragraph("4. Pipeline: operaciones en proceso con mayor probabilidad de cierre", h2)]

    datos4 = [["Proyecto", "Dpto", "Cliente", "Monto (S/)", "Prob.", "Clase"]]
    for _, r in top_pipeline.head(10).iterrows():
        datos4.append([
            str(r.get("PROYECTO", "")), str(r.get("DPTO", ""))[:8],
            str(r.get("CLIENTE", ""))[:28], f"{r.get('MONTO_TOTAL', 0):,.0f}",
            f"{r.get('PROB', 0):.0%}", str(r.get("ETIQUETA", "")),
        ])
    t4 = Table(datos4, colWidths=[2.3 * cm, 1.5 * cm, 5.4 * cm, 2.4 * cm, 1.6 * cm, 1.8 * cm])
    t4.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), AZUL),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.4, GRIS),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f3f6fa")]),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    elems += [t4, Spacer(1, 10),
              Paragraph("Umbrales de clasificación: ALTA ≥ 70 % · MEDIA 40–69 % · BAJA < 40 %. "
                        "Fuente de datos: CANASAN_2026_SIN_COLORES.xlsx (187 operaciones, 103 con resultado conocido).", body)]

    doc.build(elems)
    return buf.getvalue()
